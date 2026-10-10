#!/usr/bin/env python3
"""Generate validated government office-calendar JSON from DGPA Excel layouts.
Requires openpyxl in GitHub Actions; no guessed holidays are introduced.
"""
import argparse, calendar, datetime, json, pathlib
from openpyxl import load_workbook

PINK={'FF99FF','FFFF99FF'}
def parse(path, year):
    ws=load_workbook(path,read_only=False,data_only=True).active
    days={}; off=[]
    for month in range(1,13):
        band=(month-1)//3; col=2+((month-1)%3)*7; row=5+band*14
        for r in range(row,row+12,2):
            for dow in range(7):
                cell=ws.cell(r,col+dow); d=cell.value
                if type(d) is not int or not 1<=d<=31:continue
                try:date=datetime.date(year,month,d)
                except ValueError:continue
                if date.weekday()!=(dow+6)%7:raise ValueError(f'weekday mismatch {date}')
                key=date.isoformat()
                if key in days:raise ValueError(f'duplicate {key}')
                fg=cell.fill.fgColor
                rgb=str(fg.rgb or '').upper() if fg.type=='rgb' else ''
                holiday=rgb in PINK
                days[key]={'off':holiday,'name':'','type':'official-office-calendar'}
                if holiday:off.append(key)
    expected=366 if calendar.isleap(year) else 365
    if len(days)!=expected:raise ValueError(f'{year}: found {len(days)} dates, expected {expected}')
    if not 100<=len(off)<=130:raise ValueError(f'{year}: unusual days off {len(off)}')
    return days

def main():
    p=argparse.ArgumentParser();p.add_argument('--out',default='data/official_calendar.json');p.add_argument('--input-dir',default='official');a=p.parse_args()
    years={}
    for year in [2026,2027]:
        path=pathlib.Path(a.input_dir)/f'{year-1911}年辦公日曆表.xlsx'
        days=parse(path,year)
        years[str(year)]={'source':'DGPA Excel supplied by user','status':'official','days':days}
    labels_path=pathlib.Path(a.out).parent/'holiday_labels.json'
    labels=json.loads(labels_path.read_text(encoding='utf-8')) if labels_path.exists() else {}
    for year,rec in years.items():
        for key,name in labels.get(year,{}).items():
            if key in rec['days']:rec['days'][key]['name']=name
        rec['name_source']='Curated labels cross-checked with weekday office-calendar off dates; Excel color does not prove holiday names'
        unnamed=[key for key,day in rec['days'].items() if day['off'] and datetime.date.fromisoformat(key).weekday()<5 and not day['name']]
        if unnamed:raise ValueError(f'{year}: unnamed weekday office holidays: {unnamed}')
    result={'schema':1,'source_url':'https://www.dgpa.gov.tw/informationlist?uid=30','years':years}
    out=pathlib.Path(a.out);out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(result,ensure_ascii=False,separators=(',',':'))+'\n',encoding='utf-8')
    print('Verified:',{y:sum(x['off'] for x in rec['days'].values()) for y,rec in years.items()})
if __name__=='__main__':main()
